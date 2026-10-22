"""Ekspor tabel ke Excel, CSV, atau PDF — memakai definisi kolom yang sama dengan tampilan web."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response

from .. import auth, settings
from ..catalog import MENU_BY_KEY, Col
from ..details import DETAILS, fetch_section
from ..market import PAGE_BY_KEY, fetch_market
from ..services import dig, fetch_list
from .api import _query, _region

router = APIRouter(prefix="/api/export")

# karakter kontrol yang ditolak openpyxl (penyebab "Export Excel gagal" di versi lama)
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FORMATS = {"xlsx", "csv", "pdf"}


# ----------------------------------------------------------------------------- nilai sel
def _num(v):
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip().replace(",", ".") if v.count(",") == 1 and "." not in v else v.strip().replace(",", "")
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except ValueError:
            return v
    return v


def cell_value(col: Col, row: dict):
    v = dig(row, col.key) if col.key else None
    t = col.type
    if t == "spark":
        return None
    if v is None or v == "":
        return ""
    if t in ("num", "money"):
        return _num(v)
    if t == "tags":
        return ", ".join(str(x) for x in v) if isinstance(v, list) else str(v)
    if t == "bool":
        return "Ya" if str(v) in ("1", "True", "true") else "Tidak"
    if t == "ts":
        try:
            return dt.datetime.fromtimestamp(int(v)).strftime("%Y-%m-%d %H:%M")
        except (TypeError, ValueError, OSError):
            return str(v)
    if t == "date" and isinstance(v, (int, float)) and v > 10_000_000:
        return dt.datetime.fromtimestamp(int(v)).strftime("%Y-%m-%d")
    if t == "dur":
        try:
            s = int(float(v))
            return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"
        except (TypeError, ValueError):
            return str(v)
    if isinstance(v, (list, dict)):
        return ", ".join(map(str, v)) if isinstance(v, list) else str(v)
    return v


def table_matrix(columns: list[Col], rows: list[dict]) -> tuple[list[str], list[list]]:
    cols = [c for c in columns if c.type != "spark"]
    header = ["No"]
    for c in cols:
        header.append(c.label)
        if c.type == "entity":
            if c.sub:
                header.append(f"{c.label} — info")
            if c.id:
                header.append(f"{c.label} — ID")
    body = []
    for i, r in enumerate(rows, 1):
        line = [i]
        for c in cols:
            line.append(cell_value(c, r))
            if c.type == "entity":
                if c.sub:
                    sub = dig(r, c.sub)
                    line.append(f"{c.sub_prefix}{sub}" if sub not in (None, "") else "")
                if c.id:
                    line.append(str(dig(r, c.id) or ""))
        body.append(line)
    return header, body


def _clean(v):
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


# ----------------------------------------------------------------------------- penulis berkas
def to_xlsx(title: str, header: list, body: list, note: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "", title)[:31] or "Data"
    ws.append([_clean(title)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_clean(note)])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append([_clean(h) for h in header])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="F0365A")
        c.alignment = Alignment(vertical="center")
    for line in body:
        ws.append([_clean(v) for v in line])
    for idx, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(line[idx - 1])) for line in body[:200]])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(8, width + 2), 60)
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, int) and abs(c.value) >= 1000:
                c.number_format = "#,##0"
            elif isinstance(c.value, float):
                c.number_format = "#,##0.##"
    ws.freeze_panes = "B5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(header: list, body: list) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(body)
    return ("﻿" + buf.getvalue()).encode("utf-8")     # BOM: Excel membaca UTF-8 dengan benar


def to_pdf(title: str, header: list, body: list, note: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=6.5, leading=8)
    head = ParagraphStyle("h", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")

    def p(v, st):
        txt = _clean(str(v if v is not None else ""))
        txt = txt.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(txt[:220], st)

    data = [[p(h, head) for h in header]] + [[p(v, cell) for v in line] for line in body]
    avail = doc.width
    weights = [0.5] + [2.6 if "—" not in h and i == 0 else 1 for i, h in enumerate(header[1:])]
    total = sum(weights)
    table = LongTable(data, colWidths=[avail * w / total for w in weights], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0365A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FB")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E3E5EC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(10 * mm, 6 * mm, f"FastMoss Lite — {note}")
        canvas.drawRightString(d.pagesize[0] - 10 * mm, 6 * mm, f"Hal. {d.page}")
        canvas.restoreState()

    doc.build([Paragraph(_clean(title), ss["Title"]), Paragraph(_clean(note), ss["Normal"]), Spacer(1, 4 * mm), table],
              onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_file(fmt: str, title: str, columns: list[Col], rows: list[dict], note: str) -> Response:
    header, body = table_matrix(columns, rows)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M")
    base = re.sub(r"[^\w\- ]+", "", title).strip().replace(" ", "_") or "fastmoss"
    if fmt == "xlsx":
        content, mime = to_xlsx(title, header, body, note), \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "csv":
        content, mime = to_csv(header, body), "text/csv; charset=utf-8"
    else:
        content, mime = to_pdf(title, header, body, note), "application/pdf"
    name = f"{base}_{stamp}.{fmt}"
    return Response(content, media_type=mime,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


def _fmt(request: Request) -> str:
    fmt = (request.query_params.get("format") or "xlsx").lower()
    if fmt not in FORMATS:
        raise HTTPException(status_code=400, detail="Format harus xlsx, csv, atau pdf.")
    return fmt


def _strip(query: dict) -> dict:
    return {k: v for k, v in query.items() if k not in ("format", "page", "pages")}


def _note(region: str, rows: int, hidden: int, limit: int) -> str:
    when = dt.datetime.now().strftime("%d-%m-%Y %H:%M")
    extra = f" · {hidden}+ baris lain terkunci paket gratis (dibatasi {limit} baris)" if limit else ""
    return f"Sumber: FastMoss · Negara: {region} · {rows} baris · diunduh {when}{extra}"


def _fail(res: dict):
    status = 403 if res.get("status") in ("locked", "auth") else 502
    raise HTTPException(status_code=status, detail=res.get("message") or "Gagal mengambil data.")


# ----------------------------------------------------------------------------- endpoint
@router.get("/list/{key}")
async def export_list(key: str, request: Request, _: str = Depends(auth.require_team)):
    menu = MENU_BY_KEY.get(key)
    if not menu:
        raise HTTPException(status_code=404, detail="Menu tidak ditemukan.")
    fmt, region, query = _fmt(request), _region(request), _strip(_query(request))
    try:
        want = max(1, min(int(request.query_params.get("pages") or 1), 50))
    except ValueError:
        want = 1
    rows: list = []
    hidden = limit = 0
    for page in range(1, want + 1):
        res = await fetch_list(menu, {**query, "page": page}, region)
        if res["status"] != "ok":
            if page == 1:
                _fail(res)
            break
        rows += res["rows"]
        hidden, limit = res.get("hidden") or 0, res.get("limit") or 0
        if limit or not res["rows"] or len(rows) >= settings.EXPORT_MAX_ROWS:
            break                       # akun gratis: cukup halaman 1
        if res.get("total") and len(rows) >= res["total"]:
            break
    rows = rows[: settings.EXPORT_MAX_ROWS]
    return build_file(fmt, menu.title, menu.columns, rows, _note(region, len(rows), hidden, limit))


@router.get("/detail/{kind}/{item_id}/{section}")
async def export_detail(kind: str, item_id: str, section: str, request: Request,
                        _: str = Depends(auth.require_team)):
    det = DETAILS.get(kind)
    sec = det.by_key.get(section) if det else None
    if not sec or sec.view != "table" or not item_id.isdigit():
        raise HTTPException(status_code=404, detail="Tabel tidak ditemukan.")
    fmt, region = _fmt(request), _region(request)
    res = await fetch_section(kind, item_id, section, _strip(_query(request)), region)
    if res["status"] != "ok":
        _fail(res)
    title = f"{det.title} {item_id} - {sec.title}"
    return build_file(fmt, title, sec.fields or _auto_cols(res["rows"]), res["rows"],
                      _note(region, len(res["rows"]), res.get("hidden") or 0, res.get("limit") or 0))


@router.get("/market/{page}/{section}")
async def export_market(page: str, section: str, request: Request, _: str = Depends(auth.require_team)):
    pg = PAGE_BY_KEY.get(page)
    sec = next((s for s in pg["sections"] if s.key == section), None) if pg else None
    if not sec or sec.view != "table":
        raise HTTPException(status_code=404, detail="Tabel tidak ditemukan.")
    fmt, region = _fmt(request), _region(request)
    res = await fetch_market(page, section, _strip(_query(request)), region)
    if res["status"] != "ok":
        _fail(res)
    return build_file(fmt, f"{pg['title']} - {sec.title}", sec.fields or _auto_cols(res["rows"]), res["rows"],
                      _note(region, len(res["rows"]), res.get("hidden") or 0, res.get("limit") or 0))


def _auto_cols(rows: list[dict]) -> list[Col]:
    if not rows:
        return []
    return [Col(k, "text", k) for k, v in rows[0].items()
            if not k.endswith("_show") and not isinstance(v, (dict, list))][:30]
